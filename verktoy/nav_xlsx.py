"""Gjør en xlsx-fil fra Navs statistikkside om til TSV-filer, ett ark per fil.

Bruk: python3 verktoy/nav_xlsx.py <fil.xlsx> <utmappe> <tittel> <kildeside> [oppdatert] [neste]

Arkene hos Nav har en innledning (kilde, tittel, periode), en eller flere
tabeller med overskriftsrad (årstall, måneder eller kvartaler) og mellomoverskrifter.
Innledningen legges som #-linjer øverst i hver TSV. Tabeller i samme ark stables
med kolonnen «gruppe». Arkene «Om ...» skrives til om.md. Tall skrives som de
står i filen.
"""
import re
import sys
import unicodedata
from collections import Counter
from datetime import date
from pathlib import Path

import openpyxl


def slug(tekst):
    t = unicodedata.normalize("NFKD", tekst.lower())
    t = t.replace("æ", "ae").replace("ø", "o").replace("å", "a")
    t = "".join(c for c in t if not unicodedata.combining(c))
    return re.sub(r"[^a-z0-9]+", "-", t).strip("-")[:50]


def tekst(v):
    if v is None:
        return ""
    if isinstance(v, float):
        return str(int(v)) if v == int(v) else repr(round(v, 6))
    return re.sub(r"\s+", " ", str(v).replace("\xa0", " ")).strip()


def overskriftscelle(v):
    if isinstance(v, bool) or v is None:
        return False
    if isinstance(v, int):
        return 1990 <= v <= 2100 or (199001 <= v <= 210012 and 1 <= v % 100 <= 12)  # år eller åååmm
    s = str(v).strip()
    if re.fullmatch(r"\d{4}", s):
        return True  # årstall som tekst, f.eks. «2026» i en overskriftsrad
    return bool(s) and not set(s) <= {"*", "-", "."} and not re.fullmatch(r"[\d\s.,]+", s)


def er_overskrift(rad):
    verdier = [c for c in rad[1:] if tekst(c)]
    return len(verdier) >= 2 and all(overskriftscelle(c) for c in verdier)


def les_ark(ws):
    rader = [[(c if tekst(c) else None) for c in r] for r in ws.iter_rows(values_only=True)]
    rader = [r for r in rader if any(c is not None for c in r)]
    venstre = min(next(i for i, c in enumerate(r) if c is not None) for r in rader)
    ingress, noter, blokker, blokk, ventende, rad_gruppe = [], [], [], None, None, ""
    for r in rader:
        r = r[venstre:]
        r = r[:max(i + 1 for i, c in enumerate(r) if c is not None)]
        celler = [c for c in r if c is not None]
        if er_overskrift(r) and (len(celler) >= 3 or (len(celler) == 2 and r[0] is None)):
            blokk = {"gruppe": ventende or "", "etikett": tekst(r[0]), "kolonner": [tekst(c) for c in r[1:]], "rader": []}
            blokker.append(blokk)
            ventende, rad_gruppe = None, ""
        elif not blokker:
            ingress += [tekst(c) for c in celler]
        elif len(celler) == 1 and isinstance(celler[0], str):
            ny = tekst(celler[0])
            if len(ny) > 60 or ny.startswith(("http", "*", "-", "¹", "²", "³")):
                noter.append(ny)
            else:
                if ventende:
                    noter.append(ventende)
                ventende = ny
        else:
            if ventende:
                rad_gruppe, ventende = ventende, None
            blokk["rader"].append((rad_gruppe, [tekst(c) for c in r]))
    if ventende:
        noter.append(ventende)
    return ingress, noter, blokker


def skriv_ark(sti, ark, ingress, noter, blokker, meta):
    etiketter = {b["etikett"] for b in blokker}
    gruppert = any(b["gruppe"] or any(g for g, _ in b["rader"]) for b in blokker) or len(etiketter) > 1
    linjer = [f"# {k}: {v}" for k, v in meta]
    linjer += [f"# Ark: {ark}"] + [f"# {t}" for t in ingress if not t.startswith("Kilde")]
    linjer += [f"# Merknad: {n}" for n in noter]
    forrige = None
    delt = any(g for b in blokker for g, _ in b["rader"])
    for b in blokker:
        ar = b["rader"] and all(re.fullmatch(r"(19|20)\d\d", r[0]) for _, r in b["rader"] if r)
        navn = "år" if ar else "kategori" if re.fullmatch(r"\d\. kvartal", b["etikett"]) else b["etikett"] if len(etiketter) == 1 and b["etikett"] else "rad"
        gruppe = b["gruppe"] or (b["etikett"] or "I alt" if len(etiketter) > 1 else "")
        kol = ([("gruppe")] if gruppert else []) + [navn] + b["kolonner"]
        if kol != forrige:
            linjer.append("\t".join(kol))
            forrige = kol
        for g, r in b["rader"]:
            r = r + [""] * (len(b["kolonner"]) + 1 - len(r))
            rg = g or gruppe or ("I alt" if delt else "")
            linjer.append("\t".join(([rg] if gruppert else []) + r))
    sti.write_text("\n".join(linjer) + "\n", encoding="utf-8")


def heltall_aar(v):
    """Heltall mellom 1990 og 2100, og åååmm-koder, telles ikke: de er årstall og perioder i overskrifter og radetiketter."""
    return v == int(v) and (1990 <= v <= 2100 or (199001 <= v <= 210012 and 1 <= v % 100 <= 12))


def er_nummerrad(tall):
    """Rad med månedsnummer 1, 2, 3 ... (over månedsnavn): en overskrift, ikke data."""
    return len(tall) >= 3 and [int(t) for t in tall] == list(range(1, len(tall) + 1)) and all(t == int(t) for t in tall)


def tall_i_xlsx(wb):
    """Alle tall i datarkene, bortsett fra heltall som kan være årstall eller perioder.

    Tall som står som tekst («62», «03», «100 ») telles også, siden TSV-en ikke skiller dem fra tall.
    Rader som bare teller 1, 2, 3 ... (månedsnummer over månedsnavn) er overskrifter og telles ikke."""
    ut = Counter()
    for ws in wb:
        if ws.title.startswith("Om"):
            continue
        for rad in ws.iter_rows(values_only=True):
            celler = [c for c in rad if c is not None and tekst(c)]
            if er_nummerrad([c for c in celler if isinstance(c, (int, float)) and not isinstance(c, bool)]):
                continue
            for c in celler:
                if isinstance(c, str) and re.fullmatch(r"-?\d+([.]\d+)?", c.strip()):
                    c = float(c.strip())
                if isinstance(c, (int, float)) and not isinstance(c, bool) and not heltall_aar(c):
                    ut[round(round(float(c), 6), 3)] += 1
    return ut


def tall_i_tsv(mappe):
    ut = Counter()
    for fil in Path(mappe).glob("*.tsv"):
        for linje in fil.read_text(encoding="utf-8").splitlines():
            if linje.startswith("#"):
                continue
            tall = []
            for c in linje.split("\t"):
                try:
                    tall.append(float(c))
                except ValueError:
                    continue
            if er_nummerrad(tall):
                continue
            for v in tall:
                if not heltall_aar(v):
                    ut[round(v, 3)] += 1
    return ut


def main(fil, ut, tittel, side, oppdatert="", neste=""):
    """Konverterer fil til ut. Returnerer antall TSV-filer og tall som mangler eller er lagt til (skal være 0)."""
    ut = Path(ut)
    ut.mkdir(parents=True, exist_ok=True)
    wb = openpyxl.load_workbook(fil, data_only=True)
    meta = [("Kilde", "Nav"), ("Tabell", tittel), ("Kildeside", side)]
    meta.append(("Oppdatert", oppdatert) if oppdatert else ("Lagt inn", date.today().isoformat()))
    om = [f"# {tittel}\n", f"Kilde: Nav, {side}", (f"Oppdatert: {oppdatert}" if oppdatert else f"Lagt inn: {date.today().isoformat()}") + (f". Neste publisering: {neste}." if neste else ""), ""]
    n = 0
    for ws in wb:
        ingress, noter, blokker = les_ark(ws)
        if not blokker:
            om += [f"## {ws.title}\n"] + [f"- {t}" for t in ingress] + [""]
            continue
        skriv_ark(ut / f"{slug(ws.title)}.tsv", ws.title, ingress, noter, blokker, meta)
        n += 1
    (ut / "om.md").write_text("\n".join(om) + "\n", encoding="utf-8")
    kilde, kopi = tall_i_xlsx(wb), tall_i_tsv(ut)
    avvik = sum((kilde - kopi).values()) + sum((kopi - kilde).values())
    print(f"{n} TSV-filer i {ut}, avvik i tall: {avvik}")
    return n, avvik


if __name__ == "__main__":
    if len(sys.argv) < 5:
        sys.exit(__doc__)
    main(*sys.argv[1:])
