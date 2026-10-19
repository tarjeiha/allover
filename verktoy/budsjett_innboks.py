#!/usr/bin/env python3
"""
Legger budsjettdokumenter fra innboks-budsjett/ inn i basen.

Filnavn bestemmer hva dokumentet er (ÅR er budsjettåret, f.eks. 2027 for Prop. 1 S 2026–2027):

    prop-1-s-ÅR-aid.docx                 AIDs Prop. 1 S         -> budsjett/ÅR/prop-1-s-aid/
    gul-bok-ÅR.docx                      Prop. 1 S Gul bok      -> budsjett/ÅR/gul-bok/
    nasjonalbudsjettet-ÅR.docx           Meld. St. 1            -> budsjett/ÅR/nasjonalbudsjettet/
    prop-1-s-ÅR-tillegg-N.docx           Prop. 1 S Tillegg N    -> budsjett/ÅR/prop-1-s-tillegg-N/
    gul-bok-bevilgninger-ÅR.xlsx|csv|tsv Datagrunnlag til Gul bok -> budsjett/ÅR/gul-bok-bevilgninger.tsv

Originalnavnet fra regjeringen.no for AIDs proposisjon, prp[SESJON][NR][DEP]ddddocx.docx
(f.eks. prp202120220001asdddddocx.docx), godtas også. Budsjettåret er siste år i sesjonen.

Word-filene deles opp av verktoy/budsjett.py. Går oppdelingen bra, slettes filen fra
innboksen. Ellers blir den liggende, og en fil med samme navn og .FEIL.txt forklarer hvorfor.
Veiviseren budsjett/VEIVISER.md får en automatisk oversikt over hvilke år som ligger i basen.
"""

import csv
import re
import shutil
import sys
import tempfile
from pathlib import Path

ROT = Path(__file__).resolve().parent.parent
sys.path.insert(0, str(Path(__file__).resolve().parent))

INNBOKS = ROT / "innboks-budsjett"
BUDSJETT = ROT / "budsjett"
RAPPORT = ROT / "budsjett-rapport.txt"
START = "<!-- AUTO-START budsjettår: genereres av verktoy/budsjett_innboks.py, ikke rediger for hånd -->"
SLUTT = "<!-- AUTO-SLUTT budsjettår -->"

MONSTRE = [
    ("prop", re.compile(r"^prop-1-s-(\d{4})-aid\.docx$", re.I)),
    ("gul", re.compile(r"^gul-bok-(\d{4})\.docx$", re.I)),
    # Tilleggsproposisjon, f.eks. Støre-regjeringens endring av Solberg-regjeringens forslag for 2022.
    # Gul bok-navnet godtas også, fordi tillegget ofte lastes ned sammen med Gul bok.
    ("tillegg", re.compile(r"^(?:prop-1-s|gul-bok)-(\d{4})-tillegg[- _]?(\d{1,2})\.docx$", re.I)),
    ("nb", re.compile(r"^nasjonalbudsjettet-(\d{4})\.docx$", re.I)),
    ("data", re.compile(r"^gul-bok-bevilgninger-(\d{4})\.(xlsx|csv|tsv)$", re.I)),
]
FORVENTET_KOLONNER = {"kap_nr", "post_nr", "beløp"}
# Originalnavn fra regjeringen.no: prp + sesjon (åtte sifre) + nummer (fire sifre) + departement + ddddocx
ORIGINAL = re.compile(r"^prp(\d{4})(\d{4})(\d{4})([a-zæøå]+?)ddddocx\.docx$", re.I)
AID_KODER = {"aid", "asd"}  # AID het Arbeids- og sosialdepartementet (ASD) fram til 2022


def sjekk_aar(aar):
    if not 1990 <= aar <= 2100:
        raise ValueError(f"budsjettåret {aar} ser feil ut")
    return aar


def klassifiser(navn):
    """Gir (type, budsjettår, tilleggsnummer eller None)."""
    for type_, m in MONSTRE:
        r = m.match(navn)
        if r:
            return type_, sjekk_aar(int(r.group(1))), (int(r.group(2)) if type_ == "tillegg" else None)
    r = ORIGINAL.match(navn)
    if r:
        fra, til, nr, dep = int(r.group(1)), int(r.group(2)), int(r.group(3)), r.group(4).lower()
        if til != fra + 1:
            raise ValueError(f"sesjonen {fra}–{til} i filnavnet ser feil ut")
        if nr != 1 or dep not in AID_KODER:
            raise ValueError(
                f"filnavnet ser ut som Prop. {nr} S fra departementskoden «{dep}». Bare AIDs Prop. 1 S kjennes igjen på "
                "originalnavnet. Døp om filen etter mønsteret i README.md"
            )
        return "prop", sjekk_aar(til), None
    raise ValueError(
        "filnavnet passer ikke. Bruk prop-1-s-ÅR-aid.docx, prop-1-s-ÅR-tillegg-N.docx, gul-bok-ÅR.docx, "
        "nasjonalbudsjettet-ÅR.docx eller gul-bok-bevilgninger-ÅR.xlsx (ÅR er budsjettåret, f.eks. 2027). "
        "Originalnavnet fra regjeringen.no for AIDs Prop. 1 S (prp…ddddocx.docx) virker også"
    )


def erstatt_mappe(ny: Path, mal: Path):
    if mal.exists():
        shutil.rmtree(mal)
    mal.parent.mkdir(parents=True, exist_ok=True)
    shutil.move(str(ny), str(mal))


def del_opp_word(sti: Path, type_, aar, tillegg=None):
    import budsjett

    ut = Path(tempfile.mkdtemp()) / "ut"
    if type_ == "prop":
        filer, poster = budsjett.del_opp(str(sti), ut, sti.name)
        if len(filer) < 5 or not poster:
            raise ValueError(
                f"fant bare {len(filer)} filer og {len(poster)} poster. Stilnavnene i Word-filen svarer ikke til "
                "dem verktoy/budsjett.py forventer, så dokumentet er trolig fra en annen utgave"
            )
        erstatt_mappe(ut, BUDSJETT / str(aar) / "prop-1-s-aid")
        return f"{len(filer)} filer, {len(poster)} poster"
    filer, _ = budsjett.del_opp_generell(str(sti), ut, sti.name)
    if len(filer) < 3:
        raise ValueError(f"fant bare {len(filer)} deler. Stilnavnene i Word-filen svarer ikke til dem verktoy/budsjett.py forventer")
    mappe = {"gul": "gul-bok", "nb": "nasjonalbudsjettet", "tillegg": f"prop-1-s-tillegg-{tillegg}"}[type_]
    erstatt_mappe(ut, BUDSJETT / str(aar) / mappe)
    return f"{len(filer)} filer"


def les_rader(sti: Path):
    if sti.suffix.lower() == ".xlsx":
        import openpyxl

        wb = openpyxl.load_workbook(sti, read_only=True, data_only=True)
        # Gul bok-filene har først et ark «Forklaring», så «Data»: ta arket med de forventede kolonnene
        ark = next((a for a in wb.worksheets
                    if FORVENTET_KOLONNER <= {str(c).strip() for c in next(a.iter_rows(max_row=1, values_only=True), ())}),
                   wb.worksheets[0])
        return [["" if c is None else c for c in rad] for rad in ark.iter_rows(values_only=True)]
    tekst = sti.read_text(encoding="utf-8-sig")
    skille = max([";", "\t", ","], key=lambda s: tekst.split("\n", 1)[0].count(s))
    return [rad for rad in csv.reader(tekst.splitlines(), delimiter=skille)]


def lag_tsv(sti: Path, aar):
    rader = [r for r in les_rader(sti) if any(str(c).strip() for c in r)]
    if not rader:
        raise ValueError("filen er tom")
    overskrift = [str(c).strip() for c in rader[0]]
    mangler = FORVENTET_KOLONNER - set(overskrift)
    if mangler:
        raise ValueError(f"første rad mangler kolonnene {', '.join(sorted(mangler))}. Fant: {', '.join(overskrift[:12])}")
    mal = BUDSJETT / str(aar) / "gul-bok-bevilgninger.tsv"
    mal.parent.mkdir(parents=True, exist_ok=True)

    def celle(c):
        if isinstance(c, float) and c.is_integer():
            c = int(c)
        return str(c).replace("\t", " ").replace("\n", " ").strip()

    with open(mal, "w", encoding="utf-8", newline="") as f:
        for r in rader:
            f.write("\t".join(celle(c) for c in r) + "\n")
    return f"{len(rader) - 1} rader"


def oppdater_veiviser():
    deler = []
    for mappe in sorted((p for p in BUDSJETT.iterdir() if p.is_dir() and p.name.isdigit()), reverse=True):
        funn = []
        kjente = [("prop-1-s-aid", "Prop. 1 S AID")]
        kjente += [(p.name, f"Prop. 1 S Tillegg {p.name.rsplit('-', 1)[1]} (alle departementer)")
                   for p in sorted(mappe.glob("prop-1-s-tillegg-*")) if p.is_dir()]
        kjente += [("gul-bok", "Gul bok"), ("nasjonalbudsjettet", "Nasjonalbudsjettet"), ("stortinget", "Stortingets vedtak")]
        for navn, tekst in kjente:
            if (mappe / navn).is_dir():
                funn.append(f"{tekst} (`budsjett/{mappe.name}/{navn}/`)")
        if (mappe / "gul-bok-bevilgninger.tsv").exists():
            funn.append(f"datagrunnlag (`budsjett/{mappe.name}/gul-bok-bevilgninger.tsv`)")
        deler.append(f"- **{mappe.name}:** " + (", ".join(funn) if funn else "ingen dokumenter"))
    blokk = (
        f"{START}\n## Budsjettår i basen\n\n"
        "Hvert budsjettår har sin egen mappe. Tall (datagrunnlaget) kan gå lenger tilbake enn tekst (proposisjonene). "
        "Oppgi alltid budsjettår og dokument i svaret, og husk at Prop. 1 S er forslag, ikke vedtatt budsjett.\n\n"
        + "\n".join(deler)
        + f"\n{SLUTT}\n"
    )
    veiviser = BUDSJETT / "VEIVISER.md"
    s = veiviser.read_text(encoding="utf-8")
    if START in s and SLUTT in s:
        a, b = s.index(START), s.index(SLUTT) + len(SLUTT)
        s = s[:a] + blokk.rstrip("\n") + s[b:]
    else:
        m = re.search(r"^## Budsjettår \d{4}", s, re.M)
        if not m:
            raise ValueError("fant ikke stedet i budsjett/VEIVISER.md der oversikten skal stå")
        s = s[: m.start()] + blokk + "\n" + s[m.start():]
    veiviser.write_text(s, encoding="utf-8")


def main():
    if not INNBOKS.is_dir():
        print("Ingen innboks.")
        return 0
    filer = sorted(p for p in INNBOKS.iterdir() if p.is_file() and p.name != "README.md" and not p.name.endswith(".FEIL.txt"))
    if not filer:
        print("Innboksen er tom.")
        return 0
    ok, feil = [], []
    for sti in filer:
        feilfil = sti.with_name(sti.name + ".FEIL.txt")
        try:
            type_, aar, tillegg = klassifiser(sti.name)
            res = lag_tsv(sti, aar) if type_ == "data" else del_opp_word(sti, type_, aar, tillegg)
        except Exception as e:  # noqa: BLE001
            feilfil.write_text(f"{sti.name} ble ikke lagt inn.\n\nÅrsak: {e}\n\nRett opp og last opp filen på nytt. Denne filen forsvinner når filen er lagt inn.\n", encoding="utf-8")
            feil.append(f"{sti.name}: {e}")
            continue
        sti.unlink()
        if feilfil.exists():
            feilfil.unlink()
        ok.append(f"{sti.name}: {res}")
    if ok:
        oppdater_veiviser()
    tittel = f"Budsjett fra innboksen: {len(ok)} lagt inn, {len(feil)} stoppet"
    RAPPORT.write_text(tittel + "\n\n" + "\n".join(["Lagt inn:"] + [f"- {x}" for x in ok] + (["", "Stoppet:"] + [f"- {x}" for x in feil] if feil else [])) + "\n", encoding="utf-8")
    print(RAPPORT.read_text(encoding="utf-8"))
    return 0


if __name__ == "__main__":
    sys.exit(main())
